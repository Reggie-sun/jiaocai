"""Check current source/Word/PDF correspondence; visual review is separate."""
from pathlib import Path
from zipfile import ZipFile
import argparse
import hashlib
import json
import re

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image
from pypdf import PdfReader
import pdfplumber

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
MANUSCRIPT = ROOT / "manuscripts/ch02-progressive/ch02.md"
DOCX = ROOT / "output/ch02-progressive/第二章-AI时代计算思维的重构-递进版.docx"


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def norm(value):
    return re.sub(r"\s+|\u2060|\u200b|\*\*|`", "", value)


def inspect(pdf_path, output, toc_output=None):
    source = MANUSCRIPT.read_text()
    doc = Document(DOCX)
    table_count = len(json.loads((HERE / "assembly.json").read_text())["tables"])
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(DOCX) as archive:
        with ZipFile(ROOT / "templates/textbook-style-v1.0.dotx") as template:
            preserve = ["word/styles.xml", "word/fontTable.xml", "word/theme/theme1.xml", "word/settings.xml", "word/footer1.xml"]
            assert all(archive.read(part) == template.read(part) for part in preserve)
        xml = etree.fromstring(archive.read("word/document.xml"))
        instructions = xml.xpath("//w:instrText/text()", namespaces=ns)
        bookmarks = xml.xpath("//w:bookmarkStart/@w:name", namespaces=ns)
        assert len(bookmarks) == len(set(bookmarks))
        refs = [re.search(r"(?:PAGE)?REF\s+(\w+)", code)[1] for code in instructions if re.search(r"(?:PAGE)?REF ", code)]
        assert all(target in bookmarks for target in refs)
        assert any("TOC " in code for code in instructions)
        assert len(xml.xpath("//w:tblHeader", namespaces=ns)) == table_count
        assert len(xml.xpath("//w:tblBorders/w:left[@w:val='nil']", namespaces=ns)) == table_count
        assert all(row._tr.find(qn("w:trPr")).find(qn("w:cantSplit")) is not None for table in doc.tables for row in table.rows)
        assert len(xml.xpath("//*[local-name()='oMath']")) == 0
        assert all(f"answer_E0{i}" in bookmarks for i in range(1, 7))
        assert not any(re.search(r"(?:fig|tbl)1_", code) for code in instructions)
        word_text = "".join(xml.xpath("//w:t/text()", namespaces=ns))
    blocks, current = [], None
    for line in source.splitlines():
        if re.match(r"\*\*表2-\d+ ", line):
            current = []; blocks.append(current)
        elif line.startswith("|") and not re.match(r"^\|\s*:?-", line):
            current.append([c.strip() for c in line.strip("|").split("|")])
        elif line and not line.startswith(("#", "![", "|", "```", "$$")):
            candidate = re.sub(r"^(?:- |> )", "", line)
            assert norm(candidate) in norm(word_text), line[:100]
    actual_tables = [[ ["".join(cell._tc.xpath(".//w:t/text()")) for cell in row.cells] for row in table.rows] for table in doc.tables]
    assert [[list(map(norm, row)) for row in b] for b in blocks] == [[list(map(norm, row)) for row in b] for b in actual_tables]
    assert doc.styles["Body"].font.size.pt == 12
    assert doc.styles["Body"].paragraph_format.line_spacing == 1.5
    assert round(doc.sections[0].left_margin.mm) == 30 and round(doc.sections[0].right_margin.mm) == 25
    assert doc.settings.element.find(qn("w:mirrorMargins")) is not None
    manifest = json.loads((HERE / "figure-manifest.json").read_text())
    assert len(doc.inline_shapes) == len(manifest["figures"]) == 8
    for shape, figure in zip(doc.inline_shapes, manifest["figures"]):
        assert abs(shape.width.mm - 155) < .05
        assert sha(ROOT / figure["file"]) == figure["sha256"]
        assert sha(ROOT / figure["editable_source"]) == figure["svg_sha256"]
        with Image.open(ROOT / figure["file"]) as picture:
            assert picture.width / (shape.width.mm / 25.4) >= 300
        assert figure["min_font_pt_at_insert"] >= 10.5
    protected = json.loads((HERE / "protected-files.json").read_text())
    concurrent = json.loads((HERE / "concurrent-changes.json").read_text())["foreign_changes"]
    assert set(concurrent) == {"AGENTS.md"}
    assert all(sha(ROOT / name) == concurrent[name]["current_sha256"] if name in concurrent else sha(ROOT / name) == value for name, value in protected.items())
    reader = PdfReader(pdf_path)
    fonts = set()
    for page in reader.pages:
        for obj in page["/Resources"]["/Font"].values():
            font = obj.get_object()
            for descendant in font.get("/DescendantFonts", [obj]):
                fd = descendant.get_object().get("/FontDescriptor").get_object()
                assert any(key in fd for key in ("/FontFile", "/FontFile2", "/FontFile3"))
                fonts.add(str(fd["/FontName"]))
    texts, heading_pages, special, extents = [], {}, {}, []
    build = json.loads((HERE / "build.json").read_text())
    with pdfplumber.open(pdf_path) as pdf:
        for number, page in enumerate(pdf.pages, 1):
            text = norm("".join(c["text"] for c in page.chars)); texts.append(text)
            assert all(-.5 <= c["x0"] and c["x1"] <= page.width + .5 and -.5 <= c["top"] and c["bottom"] <= page.height + .5 for c in page.chars)
            assert not any(error in text for error in ("Error!", "错误!", "未找到引用", "Reference source not found"))
            footer = page.crop((0, page.height - 60, page.width, page.height)).extract_text().strip()
            assert footer == str(number), (number, footer)
            body = [c for c in page.chars if 70 < c["top"] < page.height - 70]
            extents.append({"page": number, "body_top": round(min(c["top"] for c in body), 1), "body_bottom": round(max(c["bottom"] for c in body), 1)})
            heading_text = norm("".join(c["text"] for c in page.chars if c["size"] >= 15))
            for heading in build["toc_headings"]:
                if heading not in heading_pages and norm(heading) in heading_text:
                    heading_pages[heading] = number
            for c in page.chars:
                if "CJK" not in c["fontname"]:
                    special.setdefault(c["fontname"], set()).add(c["text"])
    all_text = "".join(texts)
    for entry in json.loads((HERE / "assembly.json").read_text())["tables"]:
        assert norm(entry["id"] + entry["caption"]) in all_text, entry
    for index in range(1, 9):
        assert f"图2-{index}" in all_text
    assert "第二章" in all_text
    assert all(f"E0{i}参考解与讨论" in all_text for i in range(1, 7))
    assert len(heading_pages) == len(build["toc_headings"]), heading_pages
    result = {"checks": "PASS", "source_sha256": sha(MANUSCRIPT), "docx_sha256": sha(DOCX), "pdf_sha256": sha(pdf_path),
        "pages": len(reader.pages), "tables_exactly_match_source": table_count, "figures": 8, "reference_fields": len(refs),
        "toc_pages": heading_pages, "toc_cache_matches_actual_pages": build["toc_cached_pages"] == heading_pages,
        "all_fonts_embedded": sorted(fonts), "embedded_symbol_or_math_fonts": {f: "".join(sorted(chars)) for f, chars in special.items()},
        "protected_files_unchanged": len(protected) - len(concurrent), "foreign_changes_preserved": concurrent, "template_preserve_parts_identical": preserve, "page_extents": extents,
        "boundary": "Executable content/field/layout checks; does not substitute actual page viewing, classroom trial, Word native execution or publisher acceptance."}
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    if toc_output:
        toc_output.write_text(json.dumps(heading_pages, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({k: result[k] for k in ("checks", "pages", "toc_cache_matches_actual_pages", "reference_fields", "protected_files_unchanged")}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(); parser.add_argument("--pdf", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True); parser.add_argument("--toc-output", type=Path)
    arguments = parser.parse_args(); inspect(arguments.pdf, arguments.output, arguments.toc_output)
