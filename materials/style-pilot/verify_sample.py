"""Focused verification of the pilot, never a whole-chapter acceptance gate."""

import argparse
import hashlib
import importlib.util
import itertools
import json
import re
import sys
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
import pdfplumber
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
sys.dont_write_bytecode = True


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalize(text):
    return re.sub(r"\s+|\u2060", "", text)


def make_probe(sample, target):
    spec = importlib.util.spec_from_file_location("pilot", Path(__file__).with_name("build_sample.py"))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    doc = Document(sample)
    first = next(p for p in doc.paragraphs if p.style.name == "TableCaption")
    inserted = module.caption(doc, "表", 0, "插入编号测试")
    first._p.addprevious(inserted._p)
    module.numbered_heading(doc, "1.2", "同章新节测试")
    module.numbered_heading(doc, "1.2.1", "小节重置测试")
    module.chapter_title(doc, "跨章编号测试")
    module.numbered_heading(doc, "2.1", "跨章节号测试")
    module.caption(doc, "表", 99, "跨章题注探针")
    module.caption(doc, "图", 99, "跨章图号探针")
    doc.save(target)


def verify(args):
    source = ROOT / "manuscripts/style-pilot/sample.md"
    doc = Document(args.docx)
    blocks, current = [], None
    for line in source.read_text().splitlines():
        if re.match(r"表1-\d+ ", line):
            current = []
            blocks.append(current)
        elif line.startswith("|") and not re.match(r"^\|\s*:?-", line):
            current.append([cell.strip() for cell in line.strip("|").split("|")])
    assert [[list(cell.text for cell in row.cells) for row in table.rows] for table in doc.tables] == blocks
    devices = json.loads((ROOT / "materials/ch01/e01-devices.json").read_text())["devices"]
    requests = json.loads((ROOT / "materials/ch01/e01-requests.json").read_text())["requests"]
    assert blocks[0][1:] == [[d["device_id"], "display（显示）", "、".join(d["available_slots"])] for d in devices]
    assert blocks[1][1:] == [[r["request_id"], "display（显示）", "、".join(r["allowed_slots"])] for r in requests]
    expected = json.loads((ROOT / "materials/ch01/e01-reference.json").read_text())["reference"]
    assert blocks[4][1:] == [[r["request_id"], r["device_id"], r["slot_id"]] for r in expected]

    def enumerate_solutions(closed=False):
        candidates = [[(d["device_id"], s) for d in devices
                       if not (closed and d["device_id"] == "D2")
                       for s in d["available_slots"]
                       if s in r["allowed_slots"] and r["required_capability"] in d["capabilities"]]
                      for r in requests]
        return [p for p in itertools.product(*candidates) if len(set(p)) == len(requests)]

    solutions = enumerate_solutions()
    assert len(solutions) == 12 and not enumerate_solutions(True)
    assert tuple((row[1], row[2]) for row in blocks[4][1:]) in solutions
    paragraphs = "\n".join(p.text for p in doc.paragraphs)
    assert paragraphs.index("评价标准") < paragraphs.rindex("参考解与讨论")
    for line in source.read_text().splitlines():
        if line and not line.startswith(("#", "|", "![", "<!--", "表1-", "- ")):
            assert normalize(line.replace("**", "")) in normalize(paragraphs), line
    with ZipFile(args.docx) as archive:
        xml = etree.fromstring(archive.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        instructions = xml.xpath("//w:instrText/text()", namespaces=ns)
        bookmarks = xml.xpath("//w:bookmarkStart/@w:name", namespaces=ns)
        assert len(bookmarks) == len(set(bookmarks))
        assert all(re.search(r"REF\s+(\w+)", code)[1] in bookmarks
                   for code in instructions if "REF " in code)
        assert any("TOC " in code and "pilot_body" in code for code in instructions)
        assert len(xml.xpath("//w:tblHeader", namespaces=ns)) == 5
        assert all(row._tr.find(qn("w:trPr")).find(qn("w:cantSplit")) is not None
                   for table in doc.tables for row in table.rows)
        assert len(xml.xpath("//w:tblBorders/w:left[@w:val='nil']", namespaces=ns)) == 5
        numbering = etree.fromstring(archive.read("word/numbering.xml"))
        for kind in ("abstractNum", "num"):
            ids = numbering.xpath(f"//w:{kind}/@w:{kind}Id", namespaces=ns)
            assert len(ids) == len(set(ids))
    assert doc.styles["Body"].font.size.pt == 12
    assert doc.styles["Body"].paragraph_format.line_spacing == 1.5
    assert round(doc.sections[0].left_margin.mm) == 30
    assert round(doc.sections[0].right_margin.mm) == 25
    assert doc.settings.element.find(qn("w:mirrorMargins")) is not None
    with ZipFile(args.template) as archive:
        assert b"wordprocessingml.template.main+xml" in archive.read("[Content_Types].xml")
    pages = PdfReader(args.pdf).pages
    assert 4 <= len(pages) <= 8
    texts = [normalize(page.extract_text()) for page in pages]
    assert all(line.strip() != "0" for page in pages for line in page.extract_text().splitlines())
    assert "1.1从需求到可检验的过程" in texts[1]
    for n in range(1, 13):
        assert sum(text.count(f"T{n:02}") for text in texts) == 1
    assert "记录检查关注点模拟检查说明" in texts[4]
    assert "记录检查关注点模拟检查说明" in texts[5]
    fonts = set()
    for page in pages:
        for obj in page["/Resources"]["/Font"].values():
            font = obj.get_object()
            for descendant in font.get("/DescendantFonts", [obj]):
                font = descendant.get_object()
                desc = font.get("/FontDescriptor").get_object()
                assert desc and any(key in desc for key in ("/FontFile", "/FontFile2", "/FontFile3"))
                fonts.add(str(desc["/FontName"]))
    with pdfplumber.open(args.pdf) as pdf:
        for page in pdf.pages:
            assert all(c["x0"] >= 0 and c["x1"] <= page.width + .5
                       and c["top"] >= 0 and c["bottom"] <= page.height + .5 for c in page.chars)
    probe = normalize("\n".join(p.extract_text() for p in PdfReader(args.probe_pdf).pages))
    assert "表1-1插入编号测试" in probe
    assert "表1-2E01设备与开放时段" in probe
    assert "表1-2给出设备功能与开放状态，表1-3给出借用请求" in probe
    assert "表1-6E01参考安排" in probe
    assert "表2-1跨章题注探针" in probe and "2.1跨章节号测试" in probe
    assert "图2-1跨章图号探针" in probe and "1.2.1小节重置测试" in probe
    protected = json.loads(args.baseline.read_text())
    ignored = {"README.md", "教材编写现状与目的.md"}
    assert all(sha(ROOT / path) == digest for path, digest in protected.items() if path not in ignored)
    result = {"scope": "small_style_pilot", "checks": "PASS", "pages": len(pages),
              "exercise_solutions": {"baseline": 12, "D2_closed": 0},
              "protected_files_unchanged": len(protected) - len(ignored),
              "embedded_fonts": sorted(fonts), "native_numbering_insertion_and_chapter_reset": "PASS",
              "automatic_toc_cache_page": 2, "microsoft_word_execution": "NOT_EVALUATED",
              "visual_review": "separate_parent_record"}
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--make-probe", type=Path)
    parser.add_argument("--pdf", type=Path)
    parser.add_argument("--template", type=Path)
    parser.add_argument("--probe-pdf", type=Path)
    parser.add_argument("--baseline", type=Path)
    parser.add_argument("--output", type=Path)
    arguments = parser.parse_args()
    if arguments.make_probe:
        make_probe(arguments.docx, arguments.make_probe)
    else:
        verify(arguments)
