"""Build the bounded style pilot and its shared DOTX from textbook-style-v1.0.

Use the Codex bundled Python. This is a pilot adapter, not a chapter exporter.
The Skill owns the style; no chapter source or Harness file is mutated here.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SERIF = "Noto Serif CJK SC"
SANS = "Noto Sans CJK SC"
MONO = "Noto Sans Mono CJK SC"
CHAPTER = "计算思维：从现实问题到可检验的解决过程"
TABLE_WIDTHS = {
    1: (30, 57, 68), 2: (30, 57, 68), 3: (58, 43, 54),
    4: (18, 34, 103), 5: (52, 52, 51),
}


def xml(tag, **attributes):
    node = OxmlElement(tag)
    for key, value in attributes.items():
        node.set(qn("w:" + key), str(value))
    return node


def font(style, face, size, bold=False):
    style.font.name = face
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn("w:" + attr), face)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rf.attrib.pop(qn("w:" + attr), None)


def stylesheet(doc, chapter=CHAPTER):
    normal = doc.styles["Normal"]
    font(normal, SERIF, 12)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.5
    normal.element.get_or_add_pPr().append(xml("w:snapToGrid", val="false"))
    settings = doc.settings.element
    settings.append(xml("w:mirrorMargins"))
    settings.append(xml("w:updateFields", val="true"))
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(25)
    section.left_margin, section.right_margin = Mm(30), Mm(25)
    section.header_distance = section.footer_distance = Mm(15)
    grid = section._sectPr.find(qn("w:docGrid"))
    if grid is not None:
        section._sectPr.remove(grid)
    numbering = doc.part.numbering_part.element
    abstract = xml("w:abstractNum", abstractNumId="80")
    abstract.append(xml("w:multiLevelType", val="singleLevel"))
    level = xml("w:lvl", ilvl="0")
    level.append(xml("w:start", val="1"))
    level.append(xml("w:numFmt", val="chineseCounting"))
    level.append(xml("w:lvlText", val="第%1章"))
    level.append(xml("w:suff", val="space"))
    level.append(xml("w:lvlJc", val="center"))
    abstract.append(level)
    numbering.append(abstract)
    num = xml("w:num", numId="80")
    num.append(xml("w:abstractNumId", val="80"))
    numbering.append(num)

    definitions = [
        ("ChapterTitle", "Heading 1", SANS, 18, 1.5, True),
        ("SectionTitle", "Heading 2", SANS, 16, 1.5, True),
        ("SubsectionTitle", "Heading 3", SANS, 14, 1.5, True),
        ("Body", "Normal", SERIF, 12, 1.5, False),
        ("FigureCaption", "Caption", SERIF, 10.5, 1.2, False),
        ("TableCaption", "Caption", SERIF, 10.5, 1.2, False),
        ("TableBody", "Normal", SERIF, 10.5, 1.2, False),
        ("SourceNote", "Normal", SERIF, 10.5, 1.2, False),
        ("Reference", "Normal", SERIF, 10.5, 1.2, False),
        ("Code", "Normal", MONO, 10.5, 1.0, False),
    ]
    for name, base, face, size, spacing, bold in definitions:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles[base]
        font(style, face, size, bold)
        pf = style.paragraph_format
        pf.line_spacing = spacing
        pf.space_before = pf.space_after = Pt(0)
        pf.first_line_indent = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.widow_control = True
        pf.keep_with_next = False
        if name.endswith("Title"):
            pf.keep_with_next = True
            pf.keep_together = True
            level = {"ChapterTitle": 0, "SectionTitle": 1, "SubsectionTitle": 2}[name]
            style.element.get_or_add_pPr().append(xml("w:outlineLvl", val=level))
            if name == "ChapterTitle":
                pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf.page_break_before = True
                pf.space_after = Pt(18)
                num_pr = xml("w:numPr")
                num_pr.append(xml("w:ilvl", val="0"))
                num_pr.append(xml("w:numId", val="80"))
                style.element.get_or_add_pPr().append(num_pr)
            else:
                pf.space_before, pf.space_after = Pt(12), Pt(6)
        if name in ("FigureCaption", "TableCaption"):
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.space_before = pf.space_after = Pt(6)
            pf.keep_with_next = True
            pf.keep_together = True
        if name == "Body":
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.first_line_indent = Pt(24)
            pf.keep_together = True
            style.element.get_or_add_pPr().find(qn("w:ind")).set(qn("w:firstLineChars"), "200")
        if name == "Reference":
            pf.left_indent = Mm(10)
            pf.first_line_indent = Mm(-10)
            pf.tab_stops.add_tab_stop(Mm(10))
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        font(doc.styles[name], SANS, 18 if name == "Title" else 14, True)
    for name in ("Header", "Footer"):
        font(doc.styles[name], SANS, 9)
        doc.styles[name].paragraph_format.first_line_indent = Pt(0)
    for name in ("TOC 1", "TOC 2", "TOC 3"):
        style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        font(style, SERIF, 12)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.5
    hidden = doc.styles.add_style("HiddenCounter", WD_STYLE_TYPE.PARAGRAPH)
    hidden.base_style = normal
    font(hidden, SERIF, 1)
    hidden.font.hidden = True
    hidden.paragraph_format.line_spacing = Pt(1)
    hidden.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.style = "Header"
    field(header, "STYLEREF ChapterTitle", chapter)
    footer = section.footer.paragraphs[0]
    footer.style = "Footer"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(footer, "PAGE", "1")
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = "教材格式试排样张"
    doc.core_properties.subject = "textbook-style-v1.0"


def field(paragraph, instruction, cached):
    start = paragraph.add_run()
    start._r.append(xml("w:fldChar", fldCharType="begin"))
    code = paragraph.add_run()
    instr = xml("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " " + instruction + " "
    code._r.append(instr)
    paragraph.add_run()._r.append(xml("w:fldChar", fldCharType="separate"))
    paragraph.add_run(cached)
    paragraph.add_run()._r.append(xml("w:fldChar", fldCharType="end"))


def chapter_title(doc, text):
    p = doc.add_paragraph(style="ChapterTitle")
    p.add_run(text)
    # Keep the Arabic counter outside the title so STYLEREF stays readable.
    hidden_counters(doc, [("Chapter \\* ARABIC", "1")] + [
        (name + " \\r 0 \\h", "")
        for name in ("Section", "Subsection", "Figure", "Table")
    ])
    return p


def hidden_counters(doc, counters):
    counter = doc.add_paragraph(style="HiddenCounter")
    counter.paragraph_format.line_spacing = Pt(1)
    counter.paragraph_format.keep_with_next = True
    for instruction, cached in counters:
        field(counter, "SEQ " + instruction, cached)
    for run in counter.runs:
        run.font.hidden = True
        run.font.size = Pt(1)


def numbered_heading(doc, label, text):
    parts = label.split(".")
    if len(parts) == 2:
        hidden_counters(doc, [("Subsection \\r 0 \\h", "")])
    p = doc.add_paragraph(style="SectionTitle" if len(parts) == 2 else "SubsectionTitle")
    field(p, "SEQ Chapter \\c \\* ARABIC", parts[0])
    p.add_run(".")
    field(p, "SEQ Section " + ("\\s 1" if len(parts) == 2 else "\\c") + " \\* ARABIC", parts[1])
    if len(parts) == 3:
        p.add_run(".")
        field(p, "SEQ Subsection \\s 2 \\* ARABIC", parts[2])
    p.add_run(" " + text)
    return p


def inline(paragraph, text):
    pattern = r"(\*\*.*?\*\*|[图表]1-\d+|\[[12]\])"
    for token in re.split(pattern, text):
        if not token:
            continue
        if token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif re.fullmatch(r"[图表]1-\d+", token):
            name = ("fig" if token[0] == "图" else "tbl") + token[1:].replace("-", "_")
            field(paragraph, "REF " + name + " \\h", token)
        elif token in ("[1]", "[2]"):
            paragraph.add_run("\u2060")
            paragraph.add_run(token).font.superscript = True
            paragraph.add_run("\u2060")
        else:
            paragraph.add_run(token)


def caption(doc, kind, number, text):
    p = doc.add_paragraph(style="FigureCaption" if kind == "图" else "TableCaption")
    bm_id = str((100 if kind == "图" else 200) + number)
    bm_name = ("fig" if kind == "图" else "tbl") + "1_" + str(number)
    p._p.append(xml("w:bookmarkStart", id=bm_id, name=bm_name))
    p.add_run(kind)
    field(p, "SEQ Chapter \\c \\* ARABIC", "1")
    p.add_run("-")
    field(p, "SEQ " + ("Figure" if kind == "图" else "Table") + " \\s 1 \\* ARABIC", str(number))
    p._p.append(xml("w:bookmarkEnd", id=bm_id))
    p.add_run(" " + text)
    return p


def table(doc, rows, number):
    tab = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    props = tab._tbl.tblPr
    props.find(qn("w:tblW")).set(qn("w:w"), "8787")
    props.find(qn("w:tblW")).set(qn("w:type"), "dxa")
    borders = xml("w:tblBorders")
    for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
        borders.append(xml("w:" + edge, val="single" if edge in ("top", "bottom") else "nil", sz="8", color="000000"))
    props.append(borders)
    margin = xml("w:tblCellMar")
    for edge in ("top", "bottom"):
        margin.append(xml("w:" + edge, w="85", type="dxa"))
    for edge in ("left", "right"):
        margin.append(xml("w:" + edge, w="100", type="dxa"))
    props.append(margin)
    widths = TABLE_WIDTHS[number]
    for col, width in zip(tab.columns, widths):
        col.width = Mm(width)
    tab.rows[0]._tr.get_or_add_trPr().append(xml("w:tblHeader"))
    for row_index, row in enumerate(tab.rows):
        row._tr.get_or_add_trPr().append(xml("w:cantSplit"))
        for col_index, (cell, text) in enumerate(zip(row.cells, rows[row_index])):
            cell.width = Mm(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.style = "TableBody"
            inline(p, text)
            if row_index == 0:
                for run in p.runs:
                    run.bold = True
                cb = xml("w:tcBorders")
                cb.append(xml("w:bottom", val="single", sz="4", color="000000"))
                cell._tc.get_or_add_tcPr().append(cb)
            if number != 4 or row_index == len(tab.rows) - 1:
                p.paragraph_format.keep_with_next = True
    return tab


def build(out_dir, template_path):
    source = ROOT / "manuscripts/style-pilot/sample.md"
    lines = source.read_text().splitlines()
    doc = Document()
    stylesheet(doc)
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line.startswith("# "):
            chapter_title(doc, CHAPTER)
        elif line.startswith("##"):
            text = re.sub(r"^#+\s+", "", line)
            numbered = re.match(r"(1\.\d+(?:\.\d+)?)\s+(.*)", text)
            if numbered:
                if numbered[1] == "1.1":
                    doc.add_paragraph("样张目录（仅节标题）", style="SourceNote")
                    toc = doc.add_paragraph(style="SourceNote")
                    toc.paragraph_format.tab_stops.add_tab_stop(Mm(155), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
                    field(toc, 'TOC \\b pilot_body \\o "2-2" \\h \\z \\u', "1.1 从需求到可检验的过程\t2")
                heading = numbered_heading(doc, numbered[1], numbered[2])
                if numbered[1] == "1.1":
                    heading._p.insert(1, xml("w:bookmarkStart", id="500", name="pilot_body"))
            else:
                if text == "本章回顾":
                    doc.paragraphs[-1]._p.append(xml("w:bookmarkEnd", id="500"))
                doc.add_paragraph(text, style="SubsectionTitle" if line.startswith("###") else "SectionTitle")
        elif line.startswith("!["):
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.line_spacing = 1
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(ROOT / "materials/style-pilot/fig1-1.png"), width=Mm(155))
            drawing = p._p.xpath(".//wp:docPr")[0]
            drawing.set("descr", "口头需求经对象与关系整理成检查规则；偏好另行定义，失败后回查条件。")
            caption(doc, "图", 1, "需求如何变成可检查的规则")
        elif re.match(r"表1-\d+ ", line):
            match = re.match(r"表1-(\d+) (.*)", line)
            number = int(match[1])
            caption(doc, "表", number, match[2])
            rows = []
            while i < len(lines) and (not lines[i].strip() or lines[i].startswith("|")):
                row = lines[i].strip()
                i += 1
                if row.startswith("|") and not re.match(r"^\|\s*:?-", row):
                    rows.append([x.strip() for x in row.strip("|").split("|")])
            table(doc, rows, number)
        elif line.startswith("<!--"):
            continue
        elif line.startswith("来源："):
            p = doc.add_paragraph(style="SourceNote")
            inline(p, line)
        elif re.match(r"^\[\d+\] ", line):
            label, rest = line.split(" ", 1)
            p = doc.add_paragraph(style="Reference")
            p.add_run(label + "\t" + rest)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="Body")
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.left_indent = Pt(12)
            inline(p, "• " + line[2:])
        else:
            p = doc.add_paragraph(style="Body")
            inline(p, line)
    out_dir.mkdir(parents=True, exist_ok=True)
    sample_path = out_dir / "textbook-style-pilot.docx"
    doc.save(sample_path)
    template = Document()
    stylesheet(template, chapter="章名")
    chapter_title(template, "章名")
    numbered_heading(template, "1.1", "节标题")
    numbered_heading(template, "1.1.1", "小节标题")
    template.add_paragraph(style="Body")
    template.core_properties.title = "textbook-style-v1.0"
    # No diagram, answers, or mandatory per-section modules are baked into it.
    template_path.parent.mkdir(parents=True, exist_ok=True)
    raw_path = out_dir / "template-preview.docx"
    template.save(raw_path)
    with ZipFile(raw_path) as zin, ZipFile(template_path, "w", ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"wordprocessingml.document.main+xml", b"wordprocessingml.template.main+xml")
            zout.writestr(item, data)
    return sample_path, raw_path


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--template", type=Path, required=True)
    args = parser.parse_args()
    print(*build(args.output_dir, args.template), sep="\n")
