"""Verify source, template parts, native fields and the actual rendered PDF."""
from pathlib import Path
from zipfile import ZipFile
import argparse
import hashlib
import json
import re

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image
from pypdf import PdfReader
import pdfplumber

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
SOURCE = ROOT / "manuscripts/ch03-layout/ch03.md"
DOCX = ROOT / "output/ch03-layout/第三章-计算思维与人工智能素养.docx"


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def norm(value):
    return re.sub(r"\s+|\u2060|\u200b|\*\*|`", "", value)


def inspect(pdf_path, output, toc_output):
    source = SOURCE.read_text()
    assembly = json.loads((HERE / "assembly.json").read_text())
    restored = source
    for row in assembly["insertions"]:
        assert restored.count(row["inserted_text"]) == 1
        restored = restored.replace(row["inserted_text"], "", 1)
    for row in assembly["figure_status_replacements"]:
        assert restored.count(row["new"]) == 1
        restored = restored.replace(row["new"], row["old"], 1)
    assert restored.encode() == (ROOT / assembly["baseline"]).read_bytes()
    doc = Document(DOCX)
    source_urls = {value.rstrip("。；，.;") for value in re.findall(r"https?://[^\s。；，]+", source)}
    word_urls = {relation.target_ref for relation in doc.part.rels.values() if relation.reltype.endswith("/hyperlink")}
    assert source_urls == word_urls
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    inventory = json.loads((HERE / "template-parts.json").read_text())
    with ZipFile(DOCX) as archive:
        for name, item in inventory.items():
            if item["role"] == "preserve-only":
                assert hashlib.sha256(archive.read(name)).hexdigest() == item["sha256"], name
        xml = etree.fromstring(archive.read("word/document.xml"))
        instructions = xml.xpath("//w:instrText/text()", namespaces=ns)
        bookmarks = xml.xpath("//w:bookmarkStart/@w:name", namespaces=ns)
        assert len(bookmarks) == len(set(bookmarks))
        refs = [re.search(r"(?:PAGE)?REF\s+(\w+)", code)[1] for code in instructions if re.search(r"(?:PAGE)?REF ", code)]
        assert all(target in bookmarks for target in refs)
        assert any("TOC " in code for code in instructions)
        assert all(f"answer_E0{i}" in bookmarks for i in range(1, 7))
        assert not any(re.search(r"(?:fig|tbl)[12]_", code) for code in instructions)
        assert len(xml.xpath("//w:tblHeader", namespaces=ns)) == 12
        assert len(xml.xpath("//w:tblBorders/w:left[@w:val='nil']", namespaces=ns)) == 12
        assert all(row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None for table in doc.tables for row in table.rows)
        word_text = "".join(xml.xpath("//w:t/text()", namespaces=ns))
    blocks, current = [], None
    for line in source.splitlines():
        if re.match(r"\*\*表3-\d+ ", line):
            current = []; blocks.append(current)
        elif line.startswith("|") and not re.match(r"^\|\s*:?-", line):
            current.append([c.strip() for c in line.strip("|").split("|")])
        elif line and not line.startswith(("#", "![", "|", "```")):
            candidate = re.sub(r"^(?:- |> )", "", line)
            assert norm(candidate) in norm(word_text), line[:100]
    actual = [[["".join(cell._tc.xpath(".//w:t/text()")) for cell in row.cells] for row in table.rows] for table in doc.tables]
    assert [[[norm(c) for c in row] for row in b] for b in blocks] == [[[norm(c) for c in row] for row in b] for b in actual]
    assert doc.styles["Body"].font.size.pt == 12 and doc.styles["Body"].paragraph_format.line_spacing == 1.5
    assert round(doc.sections[0].left_margin.mm) == 30 and round(doc.sections[0].right_margin.mm) == 25
    assert doc.settings.element.find(qn("w:mirrorMargins")) is not None
    manifest = json.loads((HERE / "figure-manifest.json").read_text())
    assert len(doc.inline_shapes) == len(manifest["figures"]) == 5
    for shape, figure in zip(doc.inline_shapes, manifest["figures"]):
        assert abs(shape.width.mm - 155) < .05
        assert shape._inline.xpath("./wp:docPr")[0].get("descr") == figure["description"]
        relationship = shape._inline.xpath(".//a:blip")[0].get(qn("r:embed"))
        assert hashlib.sha256(doc.part.rels[relationship].target_part.blob).hexdigest() == figure["sha256"]
        assert sha(ROOT / figure["file"]) == figure["sha256"]
        assert sha(ROOT / figure["editable_source"]) == figure["svg_sha256"]
        assert sha(ROOT / figure["grayscale_png"]) == figure["gray_sha256"]
        with Image.open(ROOT / figure["file"]) as picture:
            assert picture.width / (shape.width.mm / 25.4) >= 300
        assert figure["min_font_pt_at_insert"] >= 10.5
    build = json.loads((HERE / "build.json").read_text())
    reader = PdfReader(pdf_path)
    fonts = set()
    for page in reader.pages:
        for obj in page["/Resources"]["/Font"].values():
            font = obj.get_object()
            for descendant in font.get("/DescendantFonts", [obj]):
                descriptor = descendant.get_object().get("/FontDescriptor").get_object()
                assert any(k in descriptor for k in ("/FontFile", "/FontFile2", "/FontFile3"))
                fonts.add(str(descriptor["/FontName"]))
    texts, pages, extents, figure_pages = [], {}, [], {}
    with pdfplumber.open(pdf_path) as pdf:
        for number, page in enumerate(pdf.pages, 1):
            text = norm("".join(c["text"] for c in page.chars)); texts.append(text)
            assert all(-.5 <= c["x0"] and c["x1"] <= page.width + .5 and -.5 <= c["top"] and c["bottom"] <= page.height + .5 for c in page.chars)
            assert not any(error in text for error in ("Error!", "错误!", "未找到引用", "Reference source not found"))
            footer = page.crop((0, page.height - 60, page.width, page.height)).extract_text().strip()
            assert footer == str(number), (number, footer)
            body = [c for c in page.chars if 70 < c["top"] < page.height - 70]
            extents.append({"page": number, "body_top": round(min(c["top"] for c in body), 1), "body_bottom": round(max(c["bottom"] for c in body), 1), "images": len(page.images)})
            heading_text = norm("".join(c["text"] for c in page.chars if c["size"] >= 15))
            for heading in build["toc_headings"]:
                if heading not in pages and norm(heading) in heading_text:
                    pages[heading] = number
            if page.images:
                left = (30 if number % 2 else 25) * 72 / 25.4
                right = page.width - (25 if number % 2 else 30) * 72 / 25.4
                assert all(left - .5 <= im["x0"] and im["x1"] <= right + .5
                           and im["top"] >= 25 * 72 / 25.4 - .5
                           and im["bottom"] <= page.height - 25 * 72 / 25.4 + .5
                           for im in page.images), (number, page.images)
                for figure in manifest["figures"]:
                    if norm(figure["id"] + figure["caption"]) in text:
                        figure_pages[figure["id"]] = number
    all_text = "".join(texts)
    assert "第三章" in all_text and all(f"E0{i}" in all_text for i in range(1, 7))
    assert len(pages) == len(build["toc_headings"])
    assert len(figure_pages) == 5, figure_pages
    result = {"checks": "PASS", "source_sha256": sha(SOURCE), "docx_sha256": sha(DOCX), "pdf_sha256": sha(pdf_path),
              "pages": len(reader.pages), "tables_exactly_match_source": 12, "figures": 5, "reference_fields": len(refs),
              "external_reference_hyperlinks": len(word_urls),
              "baseline_prose_exact_after_reversing_declared_changes": True,
              "figure_status_replacements": assembly["figure_status_replacements"],
              "template_preserve_parts_identical": [n for n, v in inventory.items() if v["role"] == "preserve-only"],
              "toc_pages": pages, "toc_cache_matches_actual_pages": build["toc_cached_pages"] == pages,
              "figure_pages": figure_pages, "all_fonts_embedded": sorted(fonts), "page_extents": extents,
              "boundary": "Executable correspondence and layout checks; actual page viewing is separate; Word native update, printing and classroom acceptance untested."}
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    if toc_output:
        toc_output.write_text(json.dumps(pages, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({k: result[k] for k in ("checks", "pages", "toc_cache_matches_actual_pages", "figure_pages")} ,ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--pdf", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--toc-output", type=Path)
    args = parser.parse_args()
    inspect(args.pdf, args.output, args.toc_output)
