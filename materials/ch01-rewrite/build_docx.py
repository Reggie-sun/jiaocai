"""Task-local Markdown adapter; retain the shared DOTX styles and native fields."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import argparse
import hashlib
import importlib.util
import json
import re
import sys
import tempfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Mm, Pt

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
SOURCE = ROOT / "manuscripts/ch01-rewrite/ch01.md"
TEMPLATE = ROOT / "templates/textbook-style-v1.0.dotx"
OUTPUT = ROOT / "output/ch01/第一章-计算思维.docx"
spec = importlib.util.spec_from_file_location("style_primitives", ROOT / "materials/style-pilot/build_sample.py")
style = importlib.util.module_from_spec(spec)
spec.loader.exec_module(style)
xml, field, caption = style.xml, style.field, style.caption


def hyperlink(paragraph, value):
    url = value.rstrip("。；，")
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    props.append(xml("w:color", val="000000"))
    run.append(props)
    text = OxmlElement("w:t")
    text.text = re.sub(r"([/\-_?&=])", lambda m: m[0] + "\u200b", url)
    run.append(text); link.append(run); paragraph._p.append(link)
    if url != value:
        paragraph.add_run(value[len(url):])


def inline(paragraph, value, refs=True):
    pattern = r"(\*\*.*?\*\*|`[^`]+`|https?://[^\s。；，]+|[图表]1-\d+|\[\d+\])"
    for token in re.split(pattern, value):
        if not token:
            continue
        if token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`"):
            paragraph.add_run(token[1:-1])
        elif token.startswith("http"):
            hyperlink(paragraph, token)
        elif refs and re.fullmatch(r"[图表]1-\d+", token):
            name = ("fig" if token[0] == "图" else "tbl") + token[1:].replace("-", "_")
            paragraph.add_run(token[0] + "\u00a0")
            field(paragraph, "REF " + name + " \\h", token[1:])
        elif refs and re.fullmatch(r"\[\d+\]", token):
            paragraph.add_run("\u2060")
            paragraph.add_run(token).font.superscript = True
            paragraph.add_run("\u2060")
        else:
            paragraph.add_run(token)


def bookmark(paragraph, name, identity):
    paragraph._p.insert(1, xml("w:bookmarkStart", id=str(identity), name=name))
    paragraph._p.append(xml("w:bookmarkEnd", id=str(identity)))


def toc(doc, headings, pages):
    doc.add_page_break()
    label = doc.add_paragraph("本章目录", style="SectionTitle")
    label.paragraph_format.keep_with_next = True
    first = doc.add_paragraph(style="TOC 1")
    first.add_run()._r.append(xml("w:fldChar", fldCharType="begin"))
    code = xml("w:instrText"); code.text = ' TOC \\t "SectionTitle,1" \\h \\z '
    first.add_run()._r.append(code)
    first.add_run()._r.append(xml("w:fldChar", fldCharType="separate"))
    for index, text in enumerate(headings):
        p = first if index == 0 else doc.add_paragraph(style="TOC 1")
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.tab_stops.add_tab_stop(Mm(155), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        p.add_run(text + "\t")
        field(p, "PAGEREF section_" + str(index) + " \\h", str(pages.get(text, 1)))
    doc.paragraphs[-1].add_run()._r.append(xml("w:fldChar", fldCharType="end"))
    doc.add_page_break()


def table(doc, rows, number):
    tab = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    props = tab._tbl.tblPr
    props.find(qn("w:tblW")).set(qn("w:w"), "8787")
    props.find(qn("w:tblW")).set(qn("w:type"), "dxa")
    borders = xml("w:tblBorders")
    for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
        borders.append(xml("w:" + edge, val="single" if edge in ("top", "bottom") else "nil", sz="8", color="000000"))
    props.append(borders)
    margins = xml("w:tblCellMar")
    for edge in ("top", "bottom", "left", "right"):
        margins.append(xml("w:" + edge, w="70" if edge in ("top", "bottom") else "85", type="dxa"))
    props.append(margins)
    # Allocate readable columns by current content; this is not a general layout system.
    n = len(rows[0])
    weights = [max(4, min(24, max(len(re.sub(r"\*\*|`", "", row[c])) for row in rows))) ** .5 for c in range(n)]
    widths = [155 * w / sum(weights) for w in weights]
    for column, width in zip(tab.columns, widths):
        column.width = Mm(width)
    for ri, row in enumerate(tab.rows):
        trp = row._tr.get_or_add_trPr()
        trp.append(xml("w:cantSplit"))
        if ri == 0:
            trp.append(xml("w:tblHeader"))
        for ci, cell in enumerate(row.cells):
            cell.width = Mm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]; p.style = "TableBody"
            p.paragraph_format.keep_together = False
            p.paragraph_format.keep_with_next = ri == 0 or (len(rows) <= 6 and ri < len(rows) - 1)
            inline(p, rows[ri][ci])
            if ri == 0:
                for run in p.runs:
                    run.bold = True
                cb = xml("w:tcBorders"); cb.append(xml("w:bottom", val="single", sz="4", color="000000"))
                cell._tc.get_or_add_tcPr().append(cb)
            elif re.fullmatch(r"[\d.]+", rows[ri][ci]):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return tab


def equation(doc):
    p = doc.add_paragraph(style="Body")
    p.paragraph_format.first_line_indent = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    math = OxmlElement("m:oMath")
    def mr(parent, text):
        run = OxmlElement("m:r")
        mp = OxmlElement("m:rPr"); plain = OxmlElement("m:sty"); plain.set(qn("m:val"), "p"); mp.append(plain); run.append(mp)
        rp = OxmlElement("w:rPr")
        rp.append(xml("w:rFonts", ascii="Noto Serif CJK SC", hAnsi="Noto Serif CJK SC", eastAsia="Noto Serif CJK SC", cs="Noto Serif CJK SC"))
        rp.append(xml("w:sz", val="24")); run.append(rp)
        t = OxmlElement("m:t"); t.text = text; run.append(t); parent.append(run)
    mr(math, "E = ")
    fraction = OxmlElement("m:f")
    numerator = OxmlElement("m:num"); mr(numerator, "n × P × h × d")
    denominator = OxmlElement("m:den"); mr(denominator, "1000")
    fraction.append(numerator); fraction.append(denominator); math.append(fraction)
    p._p.append(math)


def build(pages):
    with tempfile.TemporaryDirectory(prefix="ch01-template-") as temporary:
        base = Path(temporary) / "base.docx"
        with ZipFile(TEMPLATE) as source, ZipFile(base, "w", ZIP_DEFLATED) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = data.replace(b"wordprocessingml.template.main+xml", b"wordprocessingml.document.main+xml")
                target.writestr(item, data)
        doc = Document(base)
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    doc.core_properties.title = "第一章 计算思维：从现实问题到可检验的解决过程"
    doc.core_properties.subject = "textbook-style-v1.0"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    headings = [line[3:] for line in lines if line.startswith("## ") and line not in ("## 学习目标", "## 学习路线")]
    figures = {str(f["id"].split("-")[-1]): f for f in json.loads((ROOT / "assets/figures/ch01-rewrite/manifest.json").read_text())["figures"]}
    heading_index, i, code, current_figure = 0, 0, False, None
    while i < len(lines):
        line = lines[i].strip(); i += 1
        if not line:
            continue
        if line.startswith("```"):
            code = not code; continue
        if code:
            p = doc.add_paragraph(lines[i-1], style="Code")
            p.paragraph_format.keep_with_next = i < len(lines) and not lines[i].startswith("```")
            continue
        if line.startswith("# "):
            style.chapter_title(doc, line.split(" ", 2)[2]); continue
        if line.startswith("##"):
            text = re.sub(r"^#+ ", "", line)
            if text.startswith("1.1 "):
                toc(doc, headings, pages)
            numbered = re.fullmatch(r"(1\.\d+(?:\.\d+)?) (.*)", text)
            if numbered:
                p = style.numbered_heading(doc, numbered[1], numbered[2])
                if numbered[1] == "1.1":
                    p.paragraph_format.page_break_before = True
            elif line.startswith("###") and not re.match(r"E0[1-5]", text):
                p = doc.add_paragraph(style="Body"); p.paragraph_format.first_line_indent = Pt(0)
                p.add_run(text).bold = True
            else:
                p = doc.add_paragraph(text, style="SubsectionTitle" if line.startswith("###") else "SectionTitle")
            if line.startswith("## ") and text not in ("学习目标", "学习路线"):
                bookmark(p, "section_" + str(heading_index), 500 + heading_index); heading_index += 1
            continue
        if line.startswith("!["):
            match = re.fullmatch(r"!\[([^]]+)\]\(([^)]+)\)", line)
            current_figure = re.search(r"fig1-(\d+)", match[2])[1]
            p = doc.add_paragraph(style="Body")
            p.paragraph_format.first_line_indent = Pt(0); p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True; p.paragraph_format.line_spacing = 1
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(ROOT / figures[current_figure]["file"]), width=Mm(155))
            p._p.xpath(".//wp:docPr")[0].set("descr", figures[current_figure]["description"])
            continue
        match = re.fullmatch(r"(?:\*\*)?([图表])1-(\d+) (.*?)(?:\*\*)?", line)
        if match:
            kind, number, title = match[1], int(match[2]), match[3]
            p = caption(doc, kind, number, title)
            start = p._p.find(qn("w:bookmarkStart"))
            p._p.remove(start)
            p.runs[0]._r.addnext(start)
            p.paragraph_format.keep_with_next = True
            if kind == "表":
                rows = []
                while i < len(lines) and (not lines[i].strip() or lines[i].startswith("|")):
                    row = lines[i].strip(); i += 1
                    if row.startswith("|") and not re.match(r"^\|\s*:?-", row):
                        rows.append([x.strip() for x in row.strip("|").split("|")])
                table(doc, rows, number)
            continue
        if line.startswith("$$"):
            equation(doc); continue
        if line.startswith("来源："):
            p = doc.add_paragraph(style="SourceNote"); inline(p, line); continue
        if re.match(r"^\[\d+\] ", line):
            label, rest = line.split(" ", 1)
            p = doc.add_paragraph(style="Reference"); p.paragraph_format.keep_together = True
            p.add_run(label + "\t"); inline(p, rest, refs=False); continue
        p = doc.add_paragraph(style="Body")
        p.paragraph_format.keep_together = False
        if line.startswith("- ") or re.match(r"^\d+\. ", line):
            p.paragraph_format.first_line_indent = Pt(0); p.paragraph_format.left_indent = Pt(12)
            line = "• " + line[2:] if line.startswith("- ") else line
        if line.startswith("> "):
            p.paragraph_format.left_indent = Pt(12); line = line[2:]
        if line.startswith("**") and line.endswith("**"):
            p.paragraph_format.first_line_indent = Pt(0); p.paragraph_format.keep_with_next = True
        inline(p, line)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    evidence = {"template": str(TEMPLATE.relative_to(ROOT)), "template_sha256": hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(),
        "manuscript_sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(), "docx_sha256": hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
        "style": "textbook-style-v1.0", "toc_cached_pages": pages, "toc_headings": headings,
        "method": "Preserved template styles/numbering/header/footer; removed sample body, added native SEQ/REF/bookmarks/TOC/PAGEREF/PAGE and OMML fraction. No final DOCX roundtrip through LibreOffice.",
        "word_native_execution": "NOT_EVALUATED"}
    (HERE / "build.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"file": str(OUTPUT), "tables": len(doc.tables), "figures": len(doc.inline_shapes), "toc_entries": len(headings)}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(); parser.add_argument("--toc-pages", type=Path)
    args = parser.parse_args()
    build(json.loads(args.toc_pages.read_text()) if args.toc_pages else {})
