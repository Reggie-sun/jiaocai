"""Four-page comparison using the existing DOTX, actual image widths and fields."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import importlib.util
import json
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/"output/teaching-figures-pilot"
spec=importlib.util.spec_from_file_location("existing_style",ROOT/"materials/style-pilot/build_sample.py")
style=importlib.util.module_from_spec(spec)
spec.loader.exec_module(style)


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    template=ROOT/"templates/textbook-style-v1.0.dotx"
    with tempfile.TemporaryDirectory(prefix="teaching-figure-template-") as temp:
        base=Path(temp)/"base.docx"
        with ZipFile(template) as source,ZipFile(base,"w",ZIP_DEFLATED) as target:
            for name in source.namelist():
                data=source.read(name)
                if name=="[Content_Types].xml":
                    data=data.replace(b"wordprocessingml.template.main+xml",b"wordprocessingml.document.main+xml")
                target.writestr(name,data)
        doc=Document(base)
    body=doc._element.body
    for child in list(body):
        if child.tag!=qn("w:sectPr"): body.remove(child)
    contexts={
        2:"T1、T2、T3都只能选b或c，资源表还列出d。每条请求候选非空，位置总数等于请求总数，仍然无解，因为d对它们不可达。读图时，区分资源总数与这组请求可以使用的位置数。",
        6:"同一场地的占用按半开区间判断。活动A为[9:00,11:00)，与[10:00,11:00)重叠，与[11:00,12:00)相接。若清理20分钟和准备15分钟均需独占场地，则须比较扩展后的实际占用。"
    }
    captions={2:"局部候选与全局瓶颈",6:"半开区间中的连续占用、端点相接与准备清理时间"}
    for index,(number,version) in enumerate([(2,"original"),(2,"candidate"),(6,"original"),(6,"candidate")]):
        if index: doc.add_page_break()
        heading="原图" if version=="original" else "独立候选"
        doc.add_paragraph(f"图1-{number} {heading}试排",style="SectionTitle")
        p=doc.add_paragraph(contexts[number],style="Body")
        p.paragraph_format.keep_with_next=True
        name=f"{version}_fig1_{number}"
        p=doc.add_paragraph(style="Body")
        p.paragraph_format.keep_with_next=True
        p.add_run("见图")
        style.field(p,"REF "+name+" \\h",f"1-{number}")
        p.add_run("；本页与对照页采用相同正文及155 mm插图宽度。")
        image=ROOT/(f"assets/figures/ch01-rewrite/fig1-{number}.png" if version=="original" else f"assets/figures/teaching-pilot/fig1-{number}-candidate.png")
        p=doc.add_paragraph(style="Body")
        p.paragraph_format.first_line_indent=Pt(0)
        p.paragraph_format.line_spacing=1
        p.paragraph_format.keep_with_next=True
        p.paragraph_format.keep_together=True
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(image),width=Mm(155))
        p._p.xpath(".//wp:docPr")[0].set("descr",contexts[number])
        cap=doc.add_paragraph(style="FigureCaption")
        cap.paragraph_format.keep_with_next=True
        cap.add_run("图")
        cap._p.append(style.xml("w:bookmarkStart",id=str(400+index),name=name))
        style.field(cap,"SEQ Chapter \\r 1 \\* ARABIC","1")
        cap.add_run("-")
        style.field(cap,f"SEQ Figure \\r {number} \\* ARABIC",str(number))
        cap._p.append(style.xml("w:bookmarkEnd",id=str(400+index)))
        cap.add_run(" "+captions[number])
        doc.add_paragraph("来源：本书原创教学示例；基于第一章已交付内容。独立试排不替换正式稿。",style="SourceNote")
    for p in doc.sections[0].header.paragraphs:
        if p.text: p.text="第一章教学图对照试排"
    doc.save(OUT/"教学图对照试排.docx")
    (OUT/"page-map.json").write_text(json.dumps({"style":"textbook-style-v1.0","expected_pages":["fig1-2 original","fig1-2 candidate","fig1-6 original","fig1-6 candidate"],"insert_width_mm":155},indent=2)+"\n")


if __name__=="__main__": main()
