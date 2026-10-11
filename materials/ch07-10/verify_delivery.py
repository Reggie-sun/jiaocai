"""Task-local correspondence checks for the four chapter deliveries."""
from pathlib import Path
from zipfile import ZipFile
import argparse
import hashlib
import json
import re

from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image
from pypdf import PdfReader
import pdfplumber

ROOT = Path(__file__).resolve().parents[2]


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def norm(text, markdown=False):
    # Italic journal names are formatting, but code/mathematical stars are data.
    if markdown:
        text = re.sub(
            r"(`[^`]+`|\*\*.*?\*\*|(?<![\w*])\*(?!\*)[^*\n]+\*(?![\w*]))",
            lambda m: m[0][2:-2] if m[0].startswith("**") else m[0][1:-1],
            text,
        )
    return re.sub(r"\s+|\u2060|\u200b", "", text)


def numeric_citations(text):
    result = []
    for token in re.findall(r"\[(\d+(?:[-–,，]\d+)*)\]", text):
        for item in re.split(r"[,，]", token):
            bounds = re.split(r"[-–]", item)
            if len(bounds) == 1:
                result.append(int(bounds[0]))
            else:
                start, end = map(int, bounds)
                assert start <= end, token
                result.extend(range(start, end + 1))
    return result


def inspect(chapter, docx, pdf, output, toc_output):
    here = ROOT / f"materials/ch{chapter:02d}"
    build = json.loads((here / "build.json").read_text())
    source_path = ROOT / build["manuscript"]
    source = source_path.read_text()
    figures = json.loads((here / "figure-manifest.json").read_text())["figures"]
    body = re.sub(r"```.*?```", "", source, flags=re.S)
    body = re.sub(r"https?://[^\s]+", "", body)
    han = len(re.findall(r"[\u4e00-\u9fff]", body))
    assert 28000 <= han <= 40000, han
    headings = re.findall(rf"^## {chapter}\.(\d+) ", source, re.M)
    assert headings == list(map(str, range(1, {7: 5, 8: 6, 9: 5, 10: 3}[chapter] + 1))), headings
    for name in ["学习目标", "学习路线", "章末综合训练", "本章回顾", "关键术语", "参考解与讨论", "参考文献"]:
        assert f"## {name}" in source
    bib = re.findall(r"^\[(\d+)\] ", source, re.M)
    assert bib == list(map(str, range(1, len(bib) + 1))), bib
    citations = numeric_citations(source.split("## 参考文献")[0])
    cited = set(citations)
    assert cited == set(map(int, bib)), (cited, bib)
    first_use = list(dict.fromkeys(citations))
    assert first_use == list(range(1, len(bib) + 1)), first_use
    doc = Document(docx)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(docx) as z:
        inventory = json.loads((ROOT / "materials/ch03-layout/template-parts.json").read_text())
        with ZipFile(ROOT / "templates/textbook-style-v1.0.dotx") as template:
            preserved_parts = [name for name, item in inventory.items() if item["role"] == "preserve-only"]
            assert all(z.read(name) == template.read(name) for name in preserved_parts)
        xml = etree.fromstring(z.read("word/document.xml"))
        text = "".join(xml.xpath("//w:t/text()", namespaces=ns))
        bookmarks = xml.xpath("//w:bookmarkStart/@w:name", namespaces=ns)
        assert len(bookmarks) == len(set(bookmarks))
        codes = xml.xpath("//w:instrText/text()", namespaces=ns)
        refs = [re.search(r"(?:PAGE)?REF\s+(\w+)", code)[1] for code in codes if re.search(r"(?:PAGE)?REF ", code)]
        assert all(ref in bookmarks for ref in refs), set(refs) - set(bookmarks)
        assert any("TOC " in code for code in codes)
        assert not any(re.search(r"(?:fig|tbl)[123]_", code) for code in codes)
        answers = re.findall(r"^### (E\d{2}) 参考解与讨论\s*$", source, re.M)
        assert answers and all("answer_" + answer in bookmarks for answer in answers)
        for table_xml in xml.xpath("//w:tbl", namespaces=ns):
            rows_xml = table_xml.xpath("./w:tr", namespaces=ns)
            assert rows_xml[0].xpath("./w:trPr/w:tblHeader", namespaces=ns)
            assert all(row.xpath("./w:trPr/w:cantSplit", namespaces=ns) for row in rows_xml)
    grids, current, code = [], None, False
    for line in source.splitlines():
        if line.startswith("```"):
            code = not code
            continue
        if re.match(rf"(?:\*\*)?表{chapter}-\d+ ", line):
            current = []
            grids.append(current)
        elif line.startswith("|") and not re.match(r"^\|\s*:?-", line):
            assert current is not None
            current.append([c.strip() for c in line.strip("|").split("|")])
        elif line and line != ">" and not line.startswith(("#", "![", "|")):
            value = re.sub(r"^(?:- |> )", "", line)
            assert norm(value, markdown=not code) in norm(text), line[:120]
    actual = [[["".join(cell._tc.xpath(".//w:t/text()")) for cell in row.cells] for row in table.rows] for table in doc.tables]
    assert [[[norm(c, markdown=True) for c in row] for row in grid] for grid in grids] == [[[norm(c) for c in row] for row in grid] for grid in actual]
    for line in source.splitlines():
        if line.startswith("### ") and not re.match(r"### (?:\d+\.\d+|E\d{2})", line):
            matches = [p for p in doc.paragraphs if p.text == line[4:]]
            assert matches and all(p.paragraph_format.keep_with_next for p in matches), line
    assert doc.styles["Body"].font.size.pt == 12
    assert doc.styles["Body"].paragraph_format.line_spacing == 1.5
    assert round(doc.sections[0].left_margin.mm) == 30
    assert round(doc.sections[0].right_margin.mm) == 25
    assert len(doc.inline_shapes) == len(figures)
    for shape, fig in zip(doc.inline_shapes, figures):
        assert abs(shape.width.mm - 155) < .05
        assert shape._inline.xpath("./wp:docPr")[0].get("descr") == fig["description"]
        rel = shape._inline.xpath(".//a:blip")[0].get(qn("r:embed"))
        assert hashlib.sha256(doc.part.rels[rel].target_part.blob).hexdigest() == fig["sha256"]
        assert sha(ROOT / fig["file"]) == fig["sha256"]
        with Image.open(ROOT / fig["file"]) as image:
            assert image.width / (shape.width.mm / 25.4) >= 300
        assert fig["min_font_pt_at_insert"] >= 10.5
    fonts = set()
    for page in PdfReader(pdf).pages:
        for font_obj in page["/Resources"]["/Font"].values():
            font = font_obj.get_object()
            for descendant in font.get("/DescendantFonts", [font_obj]):
                descriptor = descendant.get_object()["/FontDescriptor"].get_object()
                assert any(k in descriptor for k in ["/FontFile", "/FontFile2", "/FontFile3"])
                fonts.add(str(descriptor["/FontName"]))
    pages, figure_pages, extents = {}, {}, []
    with pdfplumber.open(pdf) as document:
        for number, page in enumerate(document.pages, 1):
            text = norm("".join(c["text"] for c in page.chars))
            assert not any(marker in text for marker in ["Error!", "错误!", "未找到引用", "Referencesourcenotfound"])
            assert all(-.5 <= c["x0"] and c["x1"] <= page.width + .5 and -.5 <= c["top"] and c["bottom"] <= page.height + .5 for c in page.chars)
            footer = page.crop((0, page.height - 60, page.width, page.height)).extract_text().strip()
            assert footer == str(number), (number, footer)
            heading = norm("".join(c["text"] for c in page.chars if c["size"] >= 15))
            for title in build["toc_headings"]:
                if title not in pages and norm(title) in heading:
                    pages[title] = number
            body = [c for c in page.chars if 70 < c["top"] < page.height - 70]
            extents.append({"page": number, "body_top": round(min(c["top"] for c in body), 1) if body else None, "body_bottom": round(max(c["bottom"] for c in body), 1) if body else None, "images": len(page.images)})
            for image in page.images:
                left = (30 if number % 2 else 25) * 72 / 25.4
                right = page.width - (25 if number % 2 else 30) * 72 / 25.4
                assert left - .5 <= image["x0"] and image["x1"] <= right + .5, (number, image)
            if page.images:
                for fig in figures:
                    if norm(fig["id"] + fig["caption"]) in text:
                        figure_pages[fig["id"]] = number
        count = len(document.pages)
    assert len(pages) == len(build["toc_headings"]), (pages, build["toc_headings"])
    assert len(figure_pages) == len(figures), figure_pages
    result = {"checks": "PASS", "chapter": chapter, "han_characters_approx": han, "pages": count, "tables": len(grids), "figures": len(figures), "references": len(bib), "native_reference_fields": len(refs), "source_sha256": sha(source_path), "docx_sha256": sha(docx), "pdf_sha256": sha(pdf), "toc_pages": pages, "toc_cache_matches_actual_pages": build["toc_cached_pages"] == pages, "figure_pages": figure_pages, "embedded_fonts": sorted(fonts), "page_extents": extents, "limits": "Correspondence and executable layout checks only; actual page observation and content review recorded separately. Word native refresh, printing, teacher/student acceptance untested."}
    result.update({"preserve_only_parts_byte_identical": preserved_parts,
                   "answer_bookmarks": answers, "reference_first_use_order": first_use,
                   "table_header_repeat_and_rows_keep_together": True})
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    if toc_output:
        toc_output.write_text(json.dumps(pages, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({k: result[k] for k in ["checks", "chapter", "pages", "tables", "figures", "toc_cache_matches_actual_pages"]}))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--chapter", type=int, required=True)
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--pdf", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--toc-output", type=Path)
    args = parser.parse_args()
    inspect(args.chapter, args.docx, args.pdf, args.output, args.toc_output)
